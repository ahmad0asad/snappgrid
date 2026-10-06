"""
Script to scrub KDP/Amazon references from Word manuscripts
and export high-fidelity PDFs via Microsoft Word COM automation.
"""

import os
import sys
import docx
import win32com.client

# Ensure stdout handles utf-8
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
CLEANED_DIR = os.path.join(ROOT_DIR, 'cleaned_manuscripts')
OUTPUT_DIR = os.path.join(ROOT_DIR, 'assets', 'books')

os.makedirs(CLEANED_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

MANUSCRIPTS = [
    {
        "source": "True_Crime_Logic_Puzzle_Book_for_Adults.docx",
        "slug": "the-10-minute-detective.pdf",
        "title": "The 10-Minute Detective: True Crime Logic Puzzles for Adults"
    },
    {
        "source": "critical_thinking_puzzle_book_113p_margin_fixed.docx",
        "slug": "critical-thinking-puzzle-book.pdf",
        "title": "Critical Thinking Puzzle Book"
    },
    {
        "source": "crossword_puzzle_book_80pg.docx",
        "slug": "crossword-puzzles-for-adults.pdf",
        "title": "Crossword Puzzles for Adults: 30 Medium-to-Hard Puzzles to Sharpen the Mind"
    },
    {
        "source": "Brain_Boost_Challenge_30_Day_Brain_Training.docx",
        "slug": "brain-boost-challenge.pdf",
        "title": "Brain Boost Challenge: 30-Day Brain Training"
    },
    {
        "source": "forbidden_history.docx",
        "slug": "forbidden-history.pdf",
        "title": "Forbidden History: Shocking Trivia They Didn't Teach You in School"
    },
    {
        "source": "food_culture_trivia.docx",
        "slug": "food-culture-how-the-world-eats.pdf",
        "title": "Food, Culture & How the World Eats"
    },
    {
        "source": "unbelievable_trivia_v2.docx",
        "slug": "unbelievable-but-true-trivia.pdf",
        "title": "Unbelievable But True Trivia"
    },
    {
        "source": "trivia_challenge.docx",
        "slug": "ultimate-general-knowledge-trivia-challenge.pdf",
        "title": "The Ultimate General Knowledge Trivia Challenge"
    },
    {
        "source": "knowledge_quest_full.docx",
        "slug": "knowledge-quest-250-puzzles.pdf",
        "title": "The Knowledge Quest: 250 Mixed Puzzles Across 10 Subjects"
    }
]

def scrub_paragraph(p):
    text = p.text
    if not text:
        return False
    
    modified = False

    # Check for review request blocks
    is_review_request = (
        ("review on amazon" in text.lower()) or
        ("leave a review" in text.lower() and ("curious" in text.lower() or "independent author" in text.lower() or "journey" in text.lower())) or
        ("review request" in text.lower() and "short review helps" in text.lower()) or
        ("If these puzzles kept you up past your stop" in text)
    )

    if is_review_request:
        if len(p.runs) >= 2:
            p.runs[0].text = "FEEDBACK & CONTACT"
            p.runs[1].text = "   Enjoyed the challenge? Share your feedback at snappgrid.com/contact."
            for r in p.runs[2:]:
                r.text = ""
        elif len(p.runs) == 1:
            p.runs[0].text = "FEEDBACK & CONTACT   Enjoyed the challenge? Share your feedback at snappgrid.com/contact."
        else:
            p.text = "Enjoyed the challenge? Share your feedback at snappgrid.com/contact."
        return True

    # Check for KDP / Kindle Direct Publishing
    if "Published via Kindle Direct Publishing." in text:
        for r in p.runs:
            if "Published via Kindle Direct Publishing." in r.text:
                r.text = r.text.replace("Published via Kindle Direct Publishing.", "Published by SnappGrid · Extensive Puzzle Collection")
                modified = True
        if not modified:
            p.text = text.replace("Published via Kindle Direct Publishing.", "Published by SnappGrid · Extensive Puzzle Collection")
            modified = True
        return modified

    if "Published via Kindle Direct Publishing" in text:
        for r in p.runs:
            if "Published via Kindle Direct Publishing" in r.text:
                r.text = r.text.replace("Published via Kindle Direct Publishing", "Published by SnappGrid · Extensive Puzzle Collection")
                modified = True
        if not modified:
            p.text = text.replace("Published via Kindle Direct Publishing", "Published by SnappGrid · Extensive Puzzle Collection")
            modified = True
        return modified

    if "Kindle Direct Publishing" in text:
        for r in p.runs:
            if "Kindle Direct Publishing" in r.text:
                r.text = r.text.replace("Kindle Direct Publishing", "Published by SnappGrid · Extensive Puzzle Collection")
                modified = True
        if not modified:
            p.text = text.replace("Kindle Direct Publishing", "Published by SnappGrid · Extensive Puzzle Collection")
            modified = True
        return modified

    return False

def clean_document(src_path, out_path):
    print(f"Scrubbing: {os.path.basename(src_path)} -> {os.path.basename(out_path)}")
    doc = docx.Document(src_path)
    count = 0

    for p in doc.paragraphs:
        if scrub_paragraph(p):
            count += 1

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if scrub_paragraph(p):
                        count += 1

    for sec in doc.sections:
        for p in sec.header.paragraphs:
            if scrub_paragraph(p):
                count += 1
        for p in sec.footer.paragraphs:
            if scrub_paragraph(p):
                count += 1

    doc.save(out_path)
    print(f"  Saved cleaned document ({count} modifications)")
    return count

def export_pdfs():
    print("\n--- Phase 1: Text Scrubbing ---")
    cleaned_items = []
    for item in MANUSCRIPTS:
        src_path = os.path.join(ROOT_DIR, item["source"])
        if not os.path.exists(src_path):
            # check in books/manuscripts/
            alt_path = os.path.join(ROOT_DIR, "books", "manuscripts", item["source"])
            if os.path.exists(alt_path):
                src_path = alt_path
            else:
                raise FileNotFoundError(f"Source manuscript not found: {src_path}")

        clean_path = os.path.join(CLEANED_DIR, item["source"])
        clean_document(src_path, clean_path)
        pdf_path = os.path.join(OUTPUT_DIR, item["slug"])
        cleaned_items.append({
            "clean_docx": clean_path,
            "pdf_path": pdf_path,
            "title": item["title"],
            "slug": item["slug"]
        })

    print("\n--- Phase 2: PDF Export via Word COM Automation ---")
    wdFormatPDF = 17
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False

    try:
        for item in cleaned_items:
            docx_abs = os.path.abspath(item["clean_docx"])
            pdf_abs = os.path.abspath(item["pdf_path"])
            print(f"Converting: {os.path.basename(docx_abs)} -> {item['slug']}...")
            doc = word.Documents.Open(docx_abs, ReadOnly=True)
            doc.SaveAs2(pdf_abs, FileFormat=wdFormatPDF)
            doc.Close(SaveChanges=False)
            size_mb = os.path.getsize(pdf_abs) / (1024 * 1024)
            print(f"  ✓ Exported: {item['slug']} ({size_mb:.2f} MB)")
    finally:
        word.Quit()
        print("Word application closed.")

    print("\n--- Verification of Generated PDFs ---")
    for item in cleaned_items:
        pdf_abs = os.path.abspath(item["pdf_path"])
        if os.path.exists(pdf_abs):
            size_bytes = os.path.getsize(pdf_abs)
            size_mb = size_bytes / (1024 * 1024)
            print(f"[OK] {item['slug']}: {size_mb:.2f} MB ({size_bytes:,} bytes) at {pdf_abs}")
        else:
            print(f"[FAIL] Missing {pdf_abs}")

if __name__ == "__main__":
    export_pdfs()
